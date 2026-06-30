"""Generate the Excel workbook matching the user's exact template (6 sheets)."""
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .keywords import SKILLS, extract_skills
from .scoring import parse_salary

# ---- palette (matches the source workbook) ----
TITLE_FILL = PatternFill("solid", fgColor="2E75B6")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
NOTE_FILL = PatternFill("solid", fgColor="FCE4D6")
GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
WHITE = PatternFill("solid", fgColor="FFFFFF")

WHITE_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(color="FFFFFF", bold=True, size=13)
NOTE_FONT = Font(color="FF0000", bold=False, size=9)
BOLD = Font(bold=True, size=10)
NORMAL = Font(size=10)

CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def score_fill(score: int) -> PatternFill:
    if score >= 80:
        return GREEN
    if score >= 60:
        return YELLOW
    return RED


def post_status(job) -> str:
    s = f"Active on {job.source}"
    age = job.age_days
    if age is not None and age <= 3:
        s = f"Marked 'New' on {job.source}"
    return s


def _style_header_row(ws, row, headers, height=36):
    ws.row_dimensions[row].height = height
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row, c, text)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = CENTER_WRAP
        cell.border = BORDER


def _merged_title(ws, last_col, text):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, text)
    cell.fill = TITLE_FILL
    cell.font = TITLE_FONT
    cell.alignment = CENTER
    ws.row_dimensions[1].height = 30


def build_workbook(path, *, candidate_name, jobs, resume_skills, stats,
                   criteria, search_date=None, profile=None):
    search_date = search_date or date.today().isoformat()
    resume_skills = set(resume_skills or [])
    profile = profile or {}
    wb = Workbook()

    _sheet_listings(wb.active, candidate_name, jobs, search_date)
    _sheet_top_matches(wb.create_sheet("2. Top Matches Summary"), jobs)
    _sheet_ats(wb.create_sheet("3. ATS Keywords"), jobs, resume_skills, search_date,
               profile)
    _sheet_gap(wb.create_sheet("4. Skill Gap Analysis"), jobs, resume_skills)
    _sheet_strategy(wb.create_sheet("5. Application Strategy"), jobs)
    _sheet_summary(wb.create_sheet("6. Final Summary"), jobs, resume_skills,
                   stats, search_date)

    wb.save(path)
    return path


# ---------------- Sheet 1: Job Listings ----------------
def _sheet_listings(ws, candidate_name, jobs, search_date):
    ws.title = "1. Job Listings"
    headers = ["Match\nScore", "Role", "Company", "Location", "Salary",
               "Post Status", "Work Type", "Why It Matches", "Missing Skills",
               "Apply Link"]
    widths = {"A": 9, "B": 30, "C": 20, "D": 18, "E": 22, "F": 16, "G": 11,
              "H": 42, "I": 28, "J": 42}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    _merged_title(ws, 10,
                  f"{candidate_name.upper()} — JOB SEARCH RESULTS  |  {search_date}")
    ws.merge_cells("A2:J2")
    note = ws.cell(2, 1,
        "⚠ Note: Listings pulled live from public ATS feeds (Greenhouse, Lever, "
        "Ashby). Ghost/evergreen, scam, and stale postings have been filtered out. "
        "Verify exact post dates on the apply link before applying.")
    note.fill = NOTE_FILL
    note.font = NOTE_FONT
    note.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 27.75

    _style_header_row(ws, 3, headers)

    r = 4
    for job in jobs:
        fill = score_fill(job.match_score)
        loc_text = job.location
        if job.remote_scope:
            loc_text = f"{job.location}\n{job.remote_scope}" if job.location else job.remote_scope
        row_vals = [job.match_score, job.role, job.company, loc_text,
                    job.salary or "Not stated", post_status(job),
                    job.work_type or "—", job.why_matches, job.missing, ""]
        for c, v in enumerate(row_vals, start=1):
            cell = ws.cell(r, c, v)
            cell.fill = fill
            cell.border = BORDER
            if c == 1:
                cell.font = BOLD
                cell.alignment = CENTER
            elif c == 2:
                cell.font = BOLD
                cell.alignment = LEFT_WRAP
            else:
                cell.font = NORMAL
                cell.alignment = LEFT_WRAP
        link = ws.cell(r, 10)
        if job.apply_link:
            link.value = "Apply ↗"
            link.hyperlink = job.apply_link
            link.font = Font(color="0563C1", underline="single", size=10)
        ws.row_dimensions[r].height = 79.5
        r += 1

    ws.freeze_panes = "A4"


# ---------------- Sheet 2: Top Matches Summary ----------------
def _sheet_top_matches(ws, jobs):
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    _merged_title(ws, 2, "TOP 5 MATCHES — DETAILED ANALYSIS & APPLICATION GUIDANCE")
    _style_header_row(ws, 2, ["JOB / RANK", "ANALYSIS & APPLICATION GUIDANCE"],
                      height=20)

    medals = ["🥇", "🥈", "🥉", "#4", "#5"]
    r = 3
    for i, job in enumerate(jobs[:5]):
        tag = medals[i] if i < 3 else medals[i]
        rank = f"{tag} #{i+1} — {job.company}\n{job.role}\nScore: {job.match_score} | ATS: {job.ats_score}"
        analysis = (
            f"WHY IT FITS: {job.why_matches}\n\n"
            f"COMPENSATION: {job.salary or 'Not stated'}\n"
            f"LOCATION / TYPE: {job.location or '—'} ({job.work_type or '—'})\n"
            f"SOURCE: {job.source}\n\n"
            f"GAPS TO ADDRESS: {job.missing}\n\n"
            f"APPLY: {job.apply_link}"
        )
        a = ws.cell(r, 1, rank)
        a.fill = score_fill(job.match_score)
        a.font = BOLD
        a.alignment = TOP_WRAP
        a.border = BORDER
        b = ws.cell(r, 2, analysis)
        b.fill = WHITE
        b.font = NORMAL
        b.alignment = TOP_WRAP
        b.border = BORDER
        ws.row_dimensions[r].height = 150
        r += 1
    if not jobs:
        ws.cell(3, 1, "No qualifying matches found.").font = NORMAL


# ---------------- Sheet 3: ATS Keywords ----------------
def _sheet_ats(ws, jobs, resume_skills, search_date, profile=None):
    profile = profile or {}
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 30
    _merged_title(ws, 3, f"ATS KEYWORD ANALYSIS — {search_date}")
    _style_header_row(ws, 2, ["CATEGORY", "KEYWORDS / NOTES", "STATUS"], height=20)

    # aggregate demand across postings
    demand = Counter()
    for job in jobs:
        for sk in extract_skills(job.description + " " + job.role):
            demand[sk] += 1

    r = 3
    # candidate's own profile context from resume sections
    for category, key in (("🎓 EDUCATION — From Resume", "education"),
                          ("📜 PROFESSIONAL DEVELOPMENT — From Resume", "prof_dev")):
        for term in profile.get(key, [])[:8]:
            ws.cell(r, 1, category).alignment = LEFT_WRAP
            ws.cell(r, 2, term).alignment = LEFT_WRAP
            ws.cell(r, 3, "On Resume ✔").alignment = LEFT_WRAP
            for c in (1, 2, 3):
                ws.cell(r, c).fill = NOTE_FILL
                ws.cell(r, c).font = NORMAL
                ws.cell(r, c).border = BORDER
            r += 1

    for sk in sorted(resume_skills):
        ws.cell(r, 1, "✅ STRENGTHS — On Resume").alignment = LEFT_WRAP
        ws.cell(r, 2, sk).alignment = LEFT_WRAP
        status = "Strong ✔" + (f" — in {demand[sk]} postings" if demand.get(sk) else "")
        ws.cell(r, 3, status).alignment = LEFT_WRAP
        for c in (1, 2, 3):
            ws.cell(r, c).fill = GREEN
            ws.cell(r, c).font = NORMAL
            ws.cell(r, c).border = BORDER
        r += 1

    gaps = [sk for sk, _ in demand.most_common() if sk not in resume_skills]
    for sk in gaps:
        ws.cell(r, 1, "⚠ GAPS — Not Yet on Resume").alignment = LEFT_WRAP
        ws.cell(r, 2, sk).alignment = LEFT_WRAP
        ws.cell(r, 3, f"In {demand[sk]} of {len(jobs)} postings — add when ready").alignment = LEFT_WRAP
        for c in (1, 2, 3):
            ws.cell(r, c).fill = YELLOW
            ws.cell(r, c).font = NORMAL
            ws.cell(r, c).border = BORDER
        r += 1

    for job in jobs[:5]:
        ws.cell(r, 1, "📊 ATS MATCH ESTIMATE").alignment = LEFT_WRAP
        ws.cell(r, 2, f"{job.company} — {job.role}").alignment = LEFT_WRAP
        ws.cell(r, 3, f"~{job.ats_score} / 100").alignment = LEFT_WRAP
        for c in (1, 2, 3):
            ws.cell(r, c).font = BOLD if c == 3 else NORMAL
            ws.cell(r, c).border = BORDER
        r += 1
    if jobs:
        avg = round(sum(j.ats_score for j in jobs) / len(jobs))
        ws.cell(r, 1, "📊 ATS MATCH ESTIMATE").font = BOLD
        ws.cell(r, 2, "Avg. across today's batch").font = BOLD
        ws.cell(r, 3, f"~{avg} / 100").font = BOLD
        for c in (1, 2, 3):
            ws.cell(r, c).border = BORDER


# ---------------- Sheet 4: Skill Gap Analysis ----------------
def _sheet_gap(ws, jobs, resume_skills):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 60
    _merged_title(ws, 3, "SKILL GAP ANALYSIS — Based on Today's Active Postings")
    _style_header_row(ws, 2,
                      ["SKILL / GAP", "FREQUENCY IN POSTINGS", "RECOMMENDED ACTION"],
                      height=20)

    demand = Counter()
    for job in jobs:
        for sk in extract_skills(job.description + " " + job.role):
            if sk not in resume_skills:
                demand[sk] += 1

    total = len(jobs) or 1
    r = 3
    for sk, n in demand.most_common(15):
        ws.cell(r, 1, sk).alignment = LEFT_WRAP
        ws.cell(r, 2, f"{n} of {total} postings").alignment = LEFT_WRAP
        ws.cell(r, 3, _recommend(sk)).alignment = LEFT_WRAP
        for c in (1, 2, 3):
            ws.cell(r, c).font = NORMAL
            ws.cell(r, c).border = BORDER
            if n >= total * 0.5:
                ws.cell(r, c).fill = YELLOW
        ws.row_dimensions[r].height = 30
        r += 1
    if not demand:
        ws.cell(3, 1, "No significant gaps — resume covers the in-demand keywords.").font = NORMAL


def _recommend(skill: str) -> str:
    tips = {
        "Python": "Complete a production Python course; ship 2-3 scripted projects to GitHub.",
        "LLM Integration": "Build a small app on the Anthropic/OpenAI API; add to portfolio.",
        "RAG / Vector DB": "Complete a RAG masterclass; build a doc-Q&A demo with a vector store.",
        "Docker": "Containerize one existing project; learn compose basics.",
        "Kubernetes": "Lower priority unless targeting infra roles; learn pods/deployments.",
        "AWS": "Map your Azure experience to AWS equivalents; consider Cloud Practitioner.",
    }
    return tips.get(skill, f"Add {skill} via a focused course or a portfolio project, "
                           "then surface it on the resume.")


# ---------------- Sheet 5: Application Strategy ----------------
def _sheet_strategy(ws, jobs):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    _merged_title(ws, 2, "APPLICATION STRATEGY & ACTION PLAN")
    _style_header_row(ws, 2, ["ACTION ITEM", "DETAILS"], height=20)

    rows = []
    if jobs:
        today = jobs[:2]
        rows.append(("🎯 APPLY TODAY (Priority 1)",
                     "\n".join(f"{i+1}. {j.company} — {j.role} "
                               f"(Score {j.match_score}, ATS {j.ats_score})\n   {j.apply_link}"
                               for i, j in enumerate(today))))
        week = jobs[2:5]
        if week:
            rows.append(("📝 APPLY THIS WEEK (Priority 2)",
                         "\n".join(f"{i+3}. {j.company} — {j.role} "
                                   f"(Score {j.match_score})\n   {j.apply_link}"
                                   for i, j in enumerate(week))))
        for j in jobs[:2]:
            if j.missing and "None" not in j.missing:
                rows.append((f"✍ RESUME TIP — {j.company}",
                             f"Mirror this role's language. Surface these keywords if you "
                             f"have the experience: {j.missing}."))
    rows.append(("📅 FOLLOW-UP ACTIONS",
                 "• Tailor your resume per role using the ATS keywords tab.\n"
                 "• Apply within 24-48h of posting for best visibility.\n"
                 "• Track responses and follow up after 5-7 business days.\n"
                 "• Re-run this search every 2-3 days for fresh postings."))

    r = 3
    for item, detail in rows:
        a = ws.cell(r, 1, item)
        a.font = BOLD
        a.alignment = TOP_WRAP
        a.fill = GREEN
        a.border = BORDER
        b = ws.cell(r, 2, detail)
        b.font = NORMAL
        b.alignment = TOP_WRAP
        b.border = BORDER
        ws.row_dimensions[r].height = max(45, detail.count("\n") * 16 + 30)
        r += 1


# ---------------- Sheet 6: Final Summary ----------------
def _sheet_summary(ws, jobs, resume_skills, stats, search_date):
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 80
    _merged_title(ws, 2, f"FINAL DELIVERABLE SUMMARY  |  Search Date: {search_date}")
    _style_header_row(ws, 2, ["METRIC / FINDING", "DETAIL"], height=20)

    companies = Counter(j.company for j in jobs)
    skill_demand = Counter()
    best_pay_job = None
    best_pay = -1
    for j in jobs:
        for sk in extract_skills(j.description + " " + j.role):
            skill_demand[sk] += 1
        _, hi = parse_salary(j.salary)
        if hi and hi > best_pay:
            best_pay, best_pay_job = hi, j

    top_companies = ", ".join(c for c, _ in companies.most_common(5)) or "—"
    top_skills = ", ".join(s for s, _ in skill_demand.most_common(6)) or "—"
    best_str = (f"{best_pay_job.company} — {best_pay_job.role} "
                f"(up to ${best_pay:,})") if best_pay_job else "Not stated in postings"
    top_match = (f"{jobs[0].company} — {jobs[0].role} "
                 f"(Score {jobs[0].match_score})") if jobs else "—"

    rows = [
        ("Total Qualifying Jobs Found", f"{len(jobs)} active postings after filtering"),
        ("Raw Postings Scanned", f"{stats.get('raw', 0)} across Greenhouse, Lever, Ashby"),
        ("Filtered Out (ghost/scam/stale)", str(stats.get("dropped_quality", 0))),
        ("Filtered Out (work type)", str(stats.get("dropped_worktype", 0))),
        ("Filtered Out (location mismatch)", str(stats.get("dropped_location", 0))),
        ("Filtered Out (off-target)", str(stats.get("dropped_relevance", 0))),
        ("Top Companies Hiring", top_companies),
        ("Most Common Required Skills", top_skills),
        ("Best-Paying Role Found", best_str),
        ("Highest-Probability Match", top_match),
        ("Recommended Action This Week",
         f"Apply to {jobs[0].company} first — your strongest match." if jobs
         else "Broaden criteria or add company tokens, then re-run."),
        ("Job Boards Searched",
         "Greenhouse ✔ | Lever ✔ | Ashby ✔"
         + (" | Google Jobs ✔" if stats.get("from_google") else "")),
    ]
    if stats.get("from_google"):
        rows.insert(2, ("Postings from Google Jobs", str(stats["from_google"])))
    r = 3
    for metric, detail in rows:
        a = ws.cell(r, 1, metric)
        a.font = BOLD
        a.alignment = LEFT_WRAP
        a.fill = NOTE_FILL
        a.border = BORDER
        b = ws.cell(r, 2, detail)
        b.font = NORMAL
        b.alignment = LEFT_WRAP
        b.border = BORDER
        ws.row_dimensions[r].height = 30
        r += 1
