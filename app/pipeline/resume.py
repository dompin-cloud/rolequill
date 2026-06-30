"""Resume ingestion: extract plain text from PDF/DOCX/TXT and detect skills."""
import json
from pathlib import Path

from .keywords import extract_skills
from .profile import build_profile


def extract_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _pdf_text(path)
    if ext == ".docx":
        return _docx_text(path)
    if ext == ".txt":
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported resume type: {ext}")


def _pdf_text(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _docx_text(path: str) -> str:
    import docx
    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n".join(parts)


def analyze_resume(path: str):
    """Return (text, sorted taxonomy skills, profile dict)."""
    text = extract_text(path)
    skills = sorted(extract_skills(text))
    profile = build_profile(text)
    return text, skills, profile


def skills_to_json(skills) -> str:
    return json.dumps(list(skills))


def skills_from_json(blob) -> list[str]:
    if not blob:
        return []
    try:
        return json.loads(blob)
    except (ValueError, TypeError):
        return []
